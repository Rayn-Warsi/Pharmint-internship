"""
Backtests model_ab.py (Models A/B/C/D) and model_ensemble.py (Models F/G) at
shorter training windows, by truncating Code results/panel_dataset.csv to
t <= N months and rerunning them unmodified -- both already size their
train/val/test split off T = max(t) in the panel (see model_ab._train_window),
so feeding them a shorter panel *is* the backtest.

Restores the original (full, T=20) panel_dataset.csv when done, and leaves
Forecasted results/model_comparision.csv etc. in their normal T=20 state
since 20 is the last window run.

Writes one row per (window, model) test-set score into backtesting/backtest_log.xlsx.
"""

import ast
import io
import re
import shutil
import sys
from contextlib import redirect_stdout
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Models"))

PANEL_PATH = ROOT / "Code results" / "panel_dataset.csv"
BACKUP_PATH = ROOT / "backtesting" / "_panel_dataset_full_backup.csv"
LOG_PATH = ROOT / "backtesting" / "backtest_log.xlsx"
RAW_CACHE_PATH = ROOT / "backtesting" / "_raw_scores_cache.csv"
WINDOWS = [3, 6, 9, 12, 15, 18, 20]
MIN_T_FOR_SPLIT = 5  # train window is t=3..T-2; need T-2>=3

TEST_LINE = re.compile(r"(Model [A-Za-z][^\n]*?)\s+test \(t=(\d+)\):\s*(\{.*?\})")


def canonical_name(label):
    if label.startswith("Model A"):
        return "AR(1)" if "AR(1)" in label else ("AR(2)" if "AR(2)" in label else "Model A")
    if label.startswith("Model B"):
        return "Model B"
    if label.startswith("Model C"):
        return "Model C (PPML)"
    if label.startswith("Model D"):
        return "Model D (AB-GMM)"
    if label.startswith("Model F"):
        return "Model F (GMM+XGB residual)"
    if label.startswith("Model G"):
        return "Model G (GMM+XGB+LGBM ensemble)"
    return label


def run_and_capture(main_fn):
    buf = io.StringIO()
    try:
        with redirect_stdout(buf):
            main_fn()
        return buf.getvalue(), None
    except Exception as exc:  # noqa: BLE001 -- backtest driver, log and move on
        return buf.getvalue(), exc


def parse_rows(stdout_text, window):
    rows = []
    for label, _t, dict_text in TEST_LINE.findall(stdout_text):
        try:
            # scores mix plain floats and numpy scalars (repr'd as np.float64(...)
            # under numpy>=2) -- give eval both.
            score = eval(dict_text, {"__builtins__": {}, "np": np},
                         {"nan": float("nan"), "inf": float("inf")})
        except Exception:  # noqa: BLE001
            continue
        total_actual = score.get("total_actual")
        rows.append({
            "Training window (months)": window,
            "Model": canonical_name(label.strip()),
            "MASE": score.get("mase"),
            "RMSE (log scale)": score.get("rmse_log"),
            "RMSE (count scale)": score.get("rmse_count"),
            "Mean signed error": score.get("mean_signed_err"),
            "Total forecast error %": (
                100 * (score.get("total_pred", 0) - total_actual) / total_actual if total_actual else None
            ),
            "Date run": pd.Timestamp.now().strftime("%Y-%m-%d %H:%M"),
            "Notes": "",
        })
    return rows


def main():
    if not PANEL_PATH.exists():
        raise SystemExit(f"{PANEL_PATH} not found -- run panel_build.py first.")

    import model_ab
    import model_ensemble

    shutil.copy(PANEL_PATH, BACKUP_PATH)
    full = pd.read_csv(BACKUP_PATH)

    all_rows = []
    try:
        for window in WINDOWS:
            trunc = full[full["t"] <= window]
            T = int(trunc["t"].max())
            if T < MIN_T_FOR_SPLIT:
                print(f"window={window}: T={T} too short for the model's train/val/test split, skipping")
                continue
            trunc.to_csv(PANEL_PATH, index=False)
            print(f"=== window={window} months (T={T}) ===")

            out_ab, err_ab = run_and_capture(model_ab.main)
            rows_ab = parse_rows(out_ab, window)
            all_rows.extend(rows_ab)
            if err_ab:
                all_rows.append({"Training window (months)": window, "Model": "model_ab (rest)",
                                  "Notes": f"failed: {err_ab}"})
            print(f"  model_ab: {len(rows_ab)} rows" + (f", then failed: {err_ab}" if err_ab else ""))

            out_ens, err_ens = run_and_capture(model_ensemble.main)
            rows_ens = parse_rows(out_ens, window)
            all_rows.extend(rows_ens)
            if err_ens:
                all_rows.append({"Training window (months)": window, "Model": "model_ensemble (rest)",
                                  "Notes": f"failed: {err_ens}"})
            print(f"  model_ensemble: {len(rows_ens)} rows" + (f", then failed: {err_ens}" if err_ens else ""))
    finally:
        shutil.copy(BACKUP_PATH, PANEL_PATH)
        BACKUP_PATH.unlink(missing_ok=True)

    cols = ["Training window (months)", "Model", "MASE", "RMSE (log scale)", "RMSE (count scale)",
            "Mean signed error", "Total forecast error %", "Date run", "Notes"]
    result = pd.DataFrame(all_rows, columns=cols).sort_values(
        ["Training window (months)", "Model"], kind="stable")
    result.to_csv(RAW_CACHE_PATH, index=False)  # so a locked xlsx doesn't cost a full rerun
    try:
        write_workbook(result, LOG_PATH)
    except PermissionError:
        print(f"-> {LOG_PATH} is open elsewhere (e.g. in Excel) -- close it and rerun "
              f"`python backtesting/write_log.py` to write from the cached results "
              f"({RAW_CACHE_PATH}) without redoing the model fits.")
        return
    RAW_CACHE_PATH.unlink(missing_ok=True)
    print(f"-> {LOG_PATH} ({len(result)} rows)")


MODEL_ORDER = ["AR(1)", "AR(2)", "Model B", "Model C (PPML)", "Model D (AB-GMM)",
               "Model F (GMM+XGB residual)", "Model G (GMM+XGB+LGBM ensemble)"]


def _model_sort_key(name):
    return MODEL_ORDER.index(name) if name in MODEL_ORDER else len(MODEL_ORDER)


def _pivot_metric(df, metric):
    """One row per model, one column per training window, plus stability columns.
    CV% (std/mean) is the one that matters for 'is this model stable' -- std alone
    is meaningless when MASE and RMSE(count) sit on completely different scales."""
    scored = df.dropna(subset=[metric])
    if scored.empty:
        return pd.DataFrame()
    pivot = scored.pivot_table(index="Model", columns="Training window (months)", values=metric)
    pivot = pivot.reindex(sorted(pivot.index, key=_model_sort_key))
    pivot.columns = [f"{c}mo" for c in pivot.columns]
    window_cols = list(pivot.columns)
    pivot["Mean"] = pivot[window_cols].mean(axis=1)
    pivot["Std Dev"] = pivot[window_cols].std(axis=1)
    pivot["CV %"] = 100 * pivot["Std Dev"] / pivot["Mean"]
    pivot["Range"] = pivot[window_cols].max(axis=1) - pivot[window_cols].min(axis=1)
    return pivot.reset_index()


def _stability_summary(df):
    """One row per model: CV% for each metric, so 'which model is stable' is a
    single sort instead of eyeballing three sheets."""
    rows = []
    for model in sorted(df["Model"].dropna().unique(), key=_model_sort_key):
        sub = df[df["Model"] == model]
        row = {"Model": model, "Windows with a score": sub["MASE"].notna().sum()}
        for metric in ["MASE", "RMSE (log scale)", "RMSE (count scale)"]:
            vals = sub[metric].dropna()
            row[f"{metric} CV %"] = 100 * vals.std() / vals.mean() if len(vals) > 1 and vals.mean() else None
        rows.append(row)
    out = pd.DataFrame(rows)
    if "MASE CV %" in out.columns:
        out = out.sort_values("MASE CV %", na_position="last")
    return out


def _apply_color_scale(ws, n_data_rows, window_col_count):
    """Green = low (good/stable), red = high -- on the window-value cells only,
    so the eye sees both accuracy and swing across columns at a glance."""
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter
    if n_data_rows == 0 or window_col_count == 0:
        return
    first_col = get_column_letter(2)  # column A is "Model"
    last_col = get_column_letter(1 + window_col_count)
    rule = ColorScaleRule(start_type="min", start_color="63BE7B",
                           end_type="max", end_color="F8696B")
    ws.conditional_formatting.add(f"{first_col}2:{last_col}{n_data_rows + 1}", rule)


def write_workbook(raw, path):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _stability_summary(raw).to_excel(writer, sheet_name="Stability Summary", index=False)
        for metric, sheet in [("MASE", "MASE"), ("RMSE (log scale)", "RMSE (log)"),
                               ("RMSE (count scale)", "RMSE (count)")]:
            pivot = _pivot_metric(raw, metric)
            if pivot.empty:
                continue
            pivot.to_excel(writer, sheet_name=sheet, index=False)
            n_window_cols = sum(1 for c in pivot.columns if c.endswith("mo"))
            _apply_color_scale(writer.sheets[sheet], len(pivot), n_window_cols)
        raw.to_excel(writer, sheet_name="Raw", index=False)


if __name__ == "__main__":
    main()
